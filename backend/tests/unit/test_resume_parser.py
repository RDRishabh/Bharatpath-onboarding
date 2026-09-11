"""Local text extraction from real PDF and DOCX bytes.

The documents here are generated, not fixtures, so the test says exactly what
it put in and what it expects back. Textract will implement the same protocol;
these tests then become the comparison that shows what changing engine does to
the extracted text -- which, per invariant 1, changes the score.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pypdf
import pytest

from app.core.errors import AppError
from app.modules.resume.domain import DOCX
from app.modules.resume.parser import (
    EncryptedDocumentError,
    LegacyDocUnsupportedError,
    LocalResumeParser,
    ResumeParser,
    UnreadableDocumentError,
    UnsupportedDocumentError,
    get_resume_parser,
)

parser = LocalResumeParser()


def _pdf(pages: list[str], *, password: str | None = None) -> bytes:
    writer = pypdf.PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        writer.pages[-1]  # keep the reference explicit
        page.merge_page(_text_page(text))
    if password:
        writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _text_page(text: str) -> pypdf.PageObject:
    """A page carrying a real text-drawing operator, so extract_text sees it."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    page = pypdf.PageObject.create_blank_page(width=612, height=792)
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    content = DecodedStreamObject()
    content.set_data(stream)
    page[NameObject("/Contents")] = content
    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    resources = DictionaryObject()
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources
    return page


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- the interface --------------------------------------------------------
def test_the_local_parser_satisfies_the_protocol() -> None:
    assert isinstance(get_resume_parser(), ResumeParser)


def test_the_version_pins_every_library_that_changes_the_output() -> None:
    """Invariant 1: a score must be replayable. A pypdf upgrade changes the
    extracted text, so it must change this string -- otherwise two different
    scores look like they came from the same engine."""
    v = parser.version
    assert pypdf.__version__ in v
    assert docx.__version__ in v
    assert parser.name == "local"


# --- pdf ------------------------------------------------------------------
def test_pdf_text_is_extracted() -> None:
    result = parser.extract(content=_pdf(["Priya Sharma Senior Engineer"]), mime="application/pdf")
    assert "Priya Sharma" in result.text
    assert result.page_count == 1
    assert result.parser == "local"


def test_every_page_is_read() -> None:
    result = parser.extract(
        content=_pdf(["Page One Alpha", "Page Two Beta"]), mime="application/pdf"
    )
    assert "Alpha" in result.text and "Beta" in result.text
    assert result.page_count == 2


def test_an_encrypted_pdf_gets_its_own_code() -> None:
    """A password-protected CV is a candidate mistake with an obvious remedy,
    so it must not read as corruption."""
    with pytest.raises(EncryptedDocumentError) as exc:
        parser.extract(content=_pdf(["secret"], password="hunter2"), mime="application/pdf")
    assert exc.value.code == "resume_document_encrypted"


def test_a_corrupt_pdf_is_refused_not_crashed() -> None:
    with pytest.raises(UnreadableDocumentError):
        parser.extract(content=b"%PDF-1.7\nnot really a pdf", mime="application/pdf")


# --- docx -----------------------------------------------------------------
def test_docx_paragraphs_are_extracted() -> None:
    result = parser.extract(content=_docx_bytes(["Rahul Verma", "Backend Engineer"]), mime=DOCX)
    assert "Rahul Verma" in result.text and "Backend Engineer" in result.text


def test_docx_table_content_is_not_lost() -> None:
    """A large share of CV templates lay the whole document out in a
    borderless table. `document.paragraphs` returns nothing for those, and the
    failure is silent: an empty resume that scores as no experience at all."""
    content = _docx_bytes(
        [],
        table=[["Experience", "2019-2024"], ["Infosys", "Senior Developer"]],
    )
    result = parser.extract(content=content, mime=DOCX)
    for expected in ("Experience", "2019-2024", "Infosys", "Senior Developer"):
        assert expected in result.text, f"{expected!r} lost from a table"


def test_a_merged_heading_is_not_repeated_across_its_cells() -> None:
    result = parser.extract(content=_docx_bytes([], table=[["Skills", "Skills"]]), mime=DOCX)
    assert result.text.count("Skills") == 1


def test_a_corrupt_docx_is_refused_not_crashed() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "not xml at all <<<")
    with pytest.raises(UnreadableDocumentError):
        parser.extract(content=buf.getvalue(), mime=DOCX)


# --- unsupported ----------------------------------------------------------
def test_legacy_doc_is_refused_with_an_actionable_code() -> None:
    with pytest.raises(LegacyDocUnsupportedError) as exc:
        parser.extract(content=b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", mime="application/msword")
    assert exc.value.code == "resume_legacy_doc_unsupported"


def test_an_unknown_mime_is_refused() -> None:
    with pytest.raises(UnsupportedDocumentError):
        parser.extract(content=b"\x89PNG", mime="image/png")


def test_errors_carry_no_pre_rendered_english() -> None:
    """The candidate reads this in their own language (plan.md N9), so the
    error must carry a code and substitution params -- never a sentence."""
    for exc_type in (LegacyDocUnsupportedError, EncryptedDocumentError):
        err: AppError = exc_type()
        assert err.code and " " not in err.code
        assert isinstance(err.params, dict)
